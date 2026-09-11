from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent;O=R/'output';O.mkdir(exist_ok=True)
blocks=json.loads((R/'final_blocks.json').read_text())
pages=json.loads((R/'toc_pages.json').read_text()) if (R/'toc_pages.json').exists() else {}
doc=Document();s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
s.footer_distance=Inches(.45)
for sn in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=doc.styles[sn];st.font.name='Liberation Serif';st.font.color.rgb=RGBColor(0,0,0)
 rf=st._element.get_or_add_rPr().rFonts
 for a in list(rf.attrib):
  if 'theme' in a.lower():del rf.attrib[a]
 for p in st._element.xpath('./w:pPr/w:pBdr'):p.getparent().remove(p)
 st.paragraph_format.widow_control=True
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.04;normal.paragraph_format.space_after=Pt(5)
for sn,sz in [('Heading 1',13),('Heading 2',11.5)]:
 st=doc.styles[sn];st.font.size=Pt(sz);st.font.bold=True;st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(4);st.paragraph_format.keep_with_next=True
st=doc.styles['Title'];st.font.size=Pt(20);st.paragraph_format.space_after=Pt(6)
st=doc.styles['Subtitle'];st.font.size=Pt(12);st.paragraph_format.space_after=Pt(8)
for e in st._element.xpath('./w:rPr/w:spacing'):e.getparent().remove(e)
st=doc.styles['Caption'];st.font.size=Pt(9.5);st.font.italic=True;st.paragraph_format.space_before=Pt(6);st.paragraph_format.space_after=Pt(4);st.paragraph_format.keep_with_next=True
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
# Draw a native Word subscript for common inline mathematical symbols, not model/file identifiers.
pat=r'(?<![A-Za-z0-9_])([A-Za-zΔθ][A-Za-z]?_\{?[0-9ircabm]+\}?)(?![A-Za-z0-9_])'
def rich(p,text,math=False):
 text=text.replace('S꜀','S_c').replace('p꜀','p_c')
 for part in re.split(pat,text):
  m=re.fullmatch(r'([A-Za-zΔθ][A-Za-z]?)_\{?([0-9ircabm]+)\}?',part)
  if m:
   a=p.add_run(m[1]);a.italic=True;a=p.add_run(m[2]);a.font.subscript=True
  else:a=p.add_run(part)
  if math:a.font.name='DejaVu Serif';a.font.size=Pt(10)
 return p

def math_run(text):
 r=OxmlElement('m:r');rp=OxmlElement('m:rPr');normal=OxmlElement('m:nor');normal.set(qn('m:val'),'1');rp.append(normal);r.append(rp);t=OxmlElement('m:t');t.set(qn('xml:space'),'preserve');t.text=text;r.append(t);return r

def math_nodes(text):
 mapping=str.maketrans('₀₁₂₃₄₅₆₇₈₉ᵢᵣₘ꜀','0123456789irmc')
 # Convert Unicode subscript sequences to a canonical explicit representation.
 text=re.sub(r'([A-Za-zΣθ])([₀₁₂₃₄₅₆₇₈₉ᵢᵣₘ꜀]+)',lambda m:m[1]+'_{'+m[2].translate(mapping)+'}',text)
 pattern=r'([A-Za-zΣθΔ])([_^])(?:\{([^}]+)\}|([A-Za-z0-9]+))'
 result=[];pos=0
 for m in re.finditer(pattern,text):
  if m.start()>pos:result.append(math_run(text[pos:m.start()]))
  e=OxmlElement('m:sSub' if m[2]=='_' else 'm:sSup');base=OxmlElement('m:e');base.append(math_run(m[1]));e.append(base)
  idx=OxmlElement('m:sub' if m[2]=='_' else 'm:sup');idx.append(math_run(m[3] or m[4]));e.append(idx);result.append(e);pos=m.end()
 if pos<len(text):result.append(math_run(text[pos:]))
 return result

def add_equation(p,text):
 om=OxmlElement('m:oMath')
 for e in math_nodes(text):om.append(e)
 p._p.append(om)

def widths_for(z):
 rows=z['rows'];n=len(rows[0])
 if 'widths' in z:
  w=z['widths'];return [a*6.5/sum(w) for a in w]
 if rows[0][0]=='Code':return [.48,2.05,1.75,2.22]
 if n==7:return [1.05,1.2,.3,.55,.85,1.25,1.3]
 if n==5 and rows[0][0]=='Model':return [1.02,1.25,1.25,1.4,1.58]
 if n==5 and rows[0][0]=='R':return [.35,.7,2.15,1.65,1.65]
 if n==5:return [2.6,.4,1.3,1.2,1.0]
 if n==4 and rows[0][0]=='Quantity':return [2.5,.65,1.05,2.3]
 if n==4:return [2.3,1.35,1.35,1.5]
 if n==3 and rows[0][0]=='Memo ID':return [.9,2.35,3.25]
 if n==3:return [1.8,2.15,2.55]
 if n==2:return [1.2,5.3]
 return [6.5/n]*n

def table(z):
 rich(doc.add_paragraph(style='Caption'),f"Table {z['number']}. {z['caption']}")
 rows=z['rows'];n=len(rows[0]);w=widths_for(z)
 assert all(len(r)==n for r in rows),(z['number'],[len(r) for r in rows])
 t=doc.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,a in zip(t.columns,w):c.width=Inches(a)
 for ri,row in enumerate(rows):
  rr=t.add_row();pr=rr._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if not ri:pr.append(OxmlElement('w:tblHeader'))
  for ci,(cell,txt) in enumerate(zip(rr.cells,row)):
   cell.width=Inches(w[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
   if ri==0:txt=txt.replace('first_pass','Initial').replace('null_retry','Continue').replace('generic_retry','Generic retry').replace('verification_retry','Verification retry')
   rich(p,txt)
   if ci>0 and all(re.fullmatch(r'[-+0-9./% ]+',str(v[ci])) for v in rows[1:]):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   for run in p.runs:run.font.size=Pt(9 if n<8 else 8);run.bold=ri==0
   pr=cell._tc.get_or_add_tcPr();ma=OxmlElement('w:tcMar')
   for name,value in [('top','65'),('bottom','65'),('left','105'),('right','105')]:
    e=OxmlElement('w:'+name);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');ma.append(e)
   pr.append(ma);bd=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D0D0D0');bd.append(e)
   pr.append(bd)
   if not ri:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');pr.append(sh)

headings=[z['text'] for z in blocks if z['type']=='h1' and z.get('toc',True)]
refs=False;hindex=0
for z in blocks:
 typ=z['type'];text=z.get('text','')
 if typ in ['title','subtitle']:rich(doc.add_paragraph(style=typ.title()),text)
 elif typ=='toc':
  doc.add_page_break();doc.add_paragraph('Contents','Heading 1')
  for j,h in enumerate(headings):
   p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
   # Internal bookmarks make the Word contents navigable.
   link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'section_'+str(j))
   rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text=h;rr.append(tt);link.append(rr);p._p.append(link)
   p.add_run('\t'+str(pages.get(h,'–')))
  doc.add_page_break()
 elif typ in ['h1','h2']:
  p=rich(doc.add_paragraph(style='Heading 1' if typ=='h1' else 'Heading 2'),text)
  if typ=='h1' and z.get('toc',True):
   bm=OxmlElement('w:bookmarkStart');bm.set(qn('w:id'),str(hindex));bm.set(qn('w:name'),'section_'+str(hindex));p._p.insert(1 if p._p.pPr is not None else 0,bm)
   end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(hindex));p._p.append(end);hindex+=1
  if typ=='h1':refs=text=='References'
 elif typ=='p':
  p=doc.add_paragraph();rich(p,text)
  if text.startswith('• '):p.paragraph_format.left_indent=Inches(.14);p.paragraph_format.first_line_indent=Inches(-.14);p.paragraph_format.space_after=Pt(3)
  if refs:
   p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2)
   for run in p.runs:run.font.size=Pt(10)
  for run in p.runs:
   if '\n' in run.text:pass
 elif typ=='eq':
  p=doc.add_paragraph();add_equation(p,text);p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_together=True
 elif typ=='callout':
  p=rich(doc.add_paragraph(),text);p.paragraph_format.keep_together=True
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F2F2');p._p.get_or_add_pPr().append(sh)
  for r in p.runs:r.font.size=Pt(10)
 elif typ=='table':
  table(z)
  spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=1
  spacer.add_run().font.size=Pt(3)
 elif typ=='figure':
  p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  pic=p.add_run().add_picture(str(R/'working_source'/z['file']),width=Inches(6.5))
  pic._inline.docPr.set('descr', z['caption'])
  p=rich(doc.add_paragraph(style='Caption'),f"Figure {z['number']}. {z['caption']}");p.paragraph_format.keep_with_next=False
 else:raise ValueError(typ)
doc.save(O/'RBF_v1_33_Complete_Manuscript.docx')
# Complete editable plain-text source, with original figure links and all tables.
md=[]
for z in blocks:
 t=z['type'];tx=z.get('text','')
 if t=='toc':continue
 if t in ['title','subtitle','h1','h2']:md.append({'title':'# ','subtitle':'','h1':'## ','h2':'### '}[t]+tx)
 elif t in ['p','eq','callout']:md.append(tx)
 elif t=='table':
  md.append(f"Table {z['number']}. {z['caption']}")
  rows=z['rows'];md.append('\n'.join(['| '+' | '.join(r)+' |' for r in [rows[0],['---']*len(rows[0])]+rows[1:]]))
 elif t=='figure':md.append(f"![Figure {z['number']}. {z['caption']}]({z['file']})")
(O/'RBF_v1_33_Manuscript.md').write_text('\n\n'.join(md)+'\n')
print('Written manuscript:',len(' '.join(md).split()),'words including tables and references')
